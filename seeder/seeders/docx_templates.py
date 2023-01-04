import io
import logging

from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage

logger = logging.getLogger(__name__)
LOGO_WIDTH = Mm(90)


def replace_placeholders(template, placeholders):
    def placeholder(key, value):
        if key.endswith('_LOGO'):
            return InlineImage(doc, value, width=LOGO_WIDTH)
        else:
            return value

    try:
        doc = DocxTemplate(template)
        context = {k: placeholder(k, v) for (k, v) in placeholders.items()}
        doc.render(context)
        response_bytes = io.BytesIO()
        doc.save(response_bytes)
        return response_bytes
    except Exception as e:
        logger.warning(f'Error rendering document: {e}')
        return None

import os

from docx import Document

from seeder.seeders.docx_templates import replace_placeholders

template = f'{os.path.dirname(__file__)}/my_template.docx'


def test_text_placeholder():
    company = "HEY LAIKA"
    context = {'PLACEHOLDER': company}
    response = replace_placeholders(template, context)
    assert __contains_text(response, f'Testing {company}.')


def test_image_placeholder():
    file_name = 'ai-logo.jpg'

    context = {'PLACEHOLDER_LOGO': f'{os.path.dirname(__file__)}/{file_name}'}
    response = replace_placeholders(template, context)
    assert __contains_image(response, file_name)


def test_image_placeholder_raise_exception():
    file_name = 'ai-data-innovations-logo.png'
    context = {'PLACEHOLDER_LOGO': f'{os.path.dirname(__file__)}/{file_name}'}
    assert not replace_placeholders(template, context)


def test_missing_placeholder_as_empty():
    response = replace_placeholders(template, {})
    assert __contains_text(response, 'Testing .')


def __contains_image(file, image_name):
    image_token = f'name="{image_name}"'
    return __contains(
        file, lambda paragraph: image_token in paragraph.part.blob.decode('utf-8')
    )


def __contains_text(file, text):
    return __contains(file, lambda paragraph: text in paragraph.text)


def __contains(file, criteria):
    document = Document(file)
    for paragraph in document.paragraphs:
        if criteria(paragraph):
            return True
    return False
